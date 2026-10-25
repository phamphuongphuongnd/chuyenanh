#!/usr/bin/env python3
"""Chuan hoa test case tu bat ky format nao (md/xlsx/csv/tsv/json/yaml/txt)
thanh testcases.json theo schema chuan. Xem references/testcase-schema.md."""

import argparse
import io
import json
import re
import sys
import unicodedata
from datetime import datetime, timezone

# ---------------------------------------------------------------------------
# UTF-8 output: tranh UnicodeEncodeError khi pipe stdout tren Windows.
# ---------------------------------------------------------------------------
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
    sys.stderr.reconfigure(encoding="utf-8")

FIELDS = [
    "id", "title", "steps", "expected", "priority", "tags", "data", "depends_on",
]

# alias -> field chuan. Key da bo dau + lowercase.
ALIASES = {
    "id": "id",
    "ma": "id", "matc": "id", "stt": "id", "so": "id", "ma tc": "id",
    "test case id": "id", "case id": "id", "no": "id", "#": "id",

    "title": "title", "ten": "title", "tieu de": "title", "mo ta": "title",
    "summary": "title", "scenario": "title", "ten test case": "title",
    "ten case": "title", "test scenario": "title", "mo ta ngan": "title",

    "steps": "steps", "cac buoc": "steps", "buoc thuc hien": "steps",
    "thao tac": "steps", "step": "steps", "test steps": "steps",
    "cac buoc thuc hien": "steps", "buoc": "steps",

    "expected": "expected", "ket qua mong doi": "expected",
    "kq ky vong": "expected", "expected result": "expected",
    "ket qua ky vong": "expected", "ket qua": "expected",
    "expected results": "expected", "kq mong muon": "expected",

    "priority": "priority", "do uu tien": "priority", "muc do uu tien": "priority",
    "severity": "priority",

    "tags": "tags", "nhan": "tags", "module": "tags", "phan loai": "tags",
    "tag": "tags", "chuc nang": "tags", "khu vuc": "tags", "feature": "tags",

    "data": "data", "du lieu": "data", "input": "data", "test data": "data",
    "du lieu dau vao": "data", "data test": "data",

    "depends_on": "depends_on", "phu thuoc": "depends_on",
    "dependencies": "depends_on", "depends on": "depends_on",
    "precondition": "preconditions", "preconditions": "preconditions",
    "dieu kien tien quyet": "preconditions", "dieu kien": "preconditions",
    "tien dieu kien": "preconditions",
}

AMBIGUOUS_EXPECTED = {
    "hoat dong binh thuong", "hien thi dung", "thanh cong", "ok",
    "dung", "chay dung", "hoat dong tot", "binh thuong",
}

PRIORITY_MAP = {
    "cao": "P1", "high": "P1", "1": "P1", "critical": "P1", "p1": "P1",
    "khan cap": "P1", "urgent": "P1",
    "tb": "P2", "trung binh": "P2", "medium": "P2", "2": "P2", "p2": "P2",
    "thap": "P3", "low": "P3", "3": "P3", "p3": "P3", "minor": "P3",
}

STEP_PREFIX_RE = re.compile(r"^\s*(?:\d+[.)]|[-*•]|buoc\s*\d+\s*[:.])\s*", re.I)


def strip_accents(s):
    # "d/D" khong tach duoc bang NFD (la ky tu rieng, khong phai d + dau), xu ly rieng.
    s = s.replace("đ", "d").replace("Đ", "D")
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return s


def norm_key(s):
    s = strip_accents(str(s)).lower().strip()
    s = re.sub(r"\s+", " ", s)
    return s


def norm_text(s):
    if s is None:
        return ""
    return str(s).strip()


def detect_field(header):
    key = norm_key(header)
    if key in ALIASES:
        return ALIASES[key]
    for alias, field in ALIASES.items():
        if alias and alias in key:
            return field
    return None


def normalize_priority(raw):
    if not raw:
        return "P3"
    key = norm_key(raw)
    return PRIORITY_MAP.get(key, raw.strip().upper() if raw.strip() else "P3")


def split_steps(raw):
    if raw is None:
        return []
    if isinstance(raw, list):
        return [norm_text(s) for s in raw if norm_text(s)]
    text = str(raw)
    lines = [l for l in text.split("\n") if l.strip()]
    if len(lines) <= 1:
        # co the la 1 dong nhung co nhieu buoc noi bang so thu tu ("1. .. 2. ..")
        parts = re.split(r"(?=\d+[.)]\s)", text)
        lines = [p for p in parts if p.strip()]
    steps = []
    for l in lines:
        l = STEP_PREFIX_RE.sub("", l).strip()
        if l:
            steps.append(l)
    return steps


def split_tags(raw):
    if raw is None:
        return []
    if isinstance(raw, list):
        return [norm_text(t) for t in raw if norm_text(t)]
    text = str(raw)
    parts = re.split(r"[,;/]|\n", text)
    return [p.strip() for p in parts if p.strip()]


def parse_data(raw):
    if raw is None or raw == "":
        return {}
    if isinstance(raw, dict):
        return raw
    text = str(raw).strip()
    try:
        val = json.loads(text)
        if isinstance(val, dict):
            return val
    except (json.JSONDecodeError, TypeError):
        pass
    result = {}
    for pair in re.split(r"[,;\n]", text):
        if "=" in pair:
            k, v = pair.split("=", 1)
        elif ":" in pair:
            k, v = pair.split(":", 1)
        else:
            continue
        k = k.strip()
        if k:
            result[k] = v.strip()
    return result if result else {"raw": text}


def detect_ambiguity(case):
    flags = []
    expected = norm_text(case.get("expected"))
    if not expected:
        flags.append("thieu ket qua mong doi")
    else:
        if norm_key(expected) in AMBIGUOUS_EXPECTED:
            flags.append(f'expected mo ho: "{expected}"')
        if expected.strip().startswith("[?]"):
            flags.append("ky vong chua chot (co prefix [?])")
    title = norm_text(case.get("title"))
    if title.strip().startswith("[?]"):
        flags.append("title chua chot (co prefix [?])")
    if not case.get("steps"):
        flags.append("thieu cac buoc")
    steps_text = " ".join(case.get("steps") or []).lower()
    data = case.get("data") or {}
    if ("du lieu hop le" in steps_text or "valid data" in steps_text) and not data:
        flags.append("steps nhac du lieu hop le nhung data trong")
    return flags


def make_case(raw_fields, extra, seen_ids, counter):
    case_id = norm_text(raw_fields.get("id"))
    if not case_id:
        counter[0] += 1
        case_id = f"TC-{counter[0]:03d}"
    base_id = case_id
    suffix = "b"
    while case_id in seen_ids:
        case_id = f"{base_id}-{suffix}"
        suffix = chr(ord(suffix) + 1)
    seen_ids.add(case_id)

    case = {
        "id": case_id,
        "title": norm_text(raw_fields.get("title")),
        "priority": normalize_priority(raw_fields.get("priority")),
        "tags": split_tags(raw_fields.get("tags")),
        "preconditions": split_steps(raw_fields.get("preconditions")),
        "steps": split_steps(raw_fields.get("steps")),
        "expected": norm_text(raw_fields.get("expected")),
        "data": parse_data(raw_fields.get("data")),
        "depends_on": split_tags(raw_fields.get("depends_on")),
    }
    if extra:
        case["extra"] = extra
    flags = detect_ambiguity(case)
    case["needs_clarification"] = flags
    return case


# ---------------------------------------------------------------------------
# Parsers theo tung dinh dang input
# ---------------------------------------------------------------------------

def parse_rows(rows, headers):
    """rows: list[list[str]] voi hang dau la header (da tach rieng).
    Xu ly merge-cell: dong co id+title trong duoc noi vao case truoc do
    (KHONG fill-down cac field khac)."""
    field_map = {}
    for idx, h in enumerate(headers):
        f = detect_field(h)
        if f:
            field_map[idx] = f

    cases = []
    seen_ids = set()
    counter = [0]
    current = None

    for row in rows:
        raw = {}
        extra = {}
        for idx, cell in enumerate(row):
            if idx >= len(headers):
                continue
            val = cell
            field = field_map.get(idx)
            if field:
                if field in raw and raw[field]:
                    raw[field] = f"{raw[field]}\n{val}" if val else raw[field]
                else:
                    raw[field] = val
            elif norm_text(val):
                extra[headers[idx]] = val

        is_continuation = (
            current is not None
            and not norm_text(raw.get("id"))
            and not norm_text(raw.get("title"))
            and any(norm_text(v) for v in raw.values())
        )
        if is_continuation:
            # noi them steps/expected vao case truoc, KHONG fill-down cac field khac
            for f in ("steps", "expected"):
                v = norm_text(raw.get(f))
                if not v:
                    continue
                if f == "steps":
                    current["steps"].extend(split_steps(v))
                else:
                    current["expected"] = (current["expected"] + "\n" + v).strip()
            current["needs_clarification"] = detect_ambiguity(current)
            continue

        if not any(norm_text(v) for v in raw.values()):
            continue  # dong trong hoan toan

        case = make_case(raw, extra, seen_ids, counter)
        cases.append(case)
        current = case  # dong tiep theo (neu la continuation) se noi vao case nay

    return cases


def parse_markdown_table(text):
    lines = [l for l in text.splitlines() if l.strip().startswith("|")]
    if not lines:
        return [], []
    def split_row(line):
        line = line.strip()
        if line.startswith("|"):
            line = line[1:]
        if line.endswith("|"):
            line = line[:-1]
        return [c.strip() for c in line.split("|")]

    headers = split_row(lines[0])
    body = lines[1:]
    if len(body) > 1 and re.match(r"^[\s:|-]+$", body[0]):
        body = body[1:]
    rows = [split_row(l) for l in body]
    return headers, rows


def parse_csv_like(text, delimiter=","):
    import csv
    reader = csv.reader(io.StringIO(text), delimiter=delimiter)
    rows = [r for r in reader if any(c.strip() for c in r)]
    if not rows:
        return [], []
    return rows[0], rows[1:]


def parse_xlsx(path, sheet=None):
    try:
        import openpyxl
    except ImportError:
        print("Loi: doc file .xlsx can cai openpyxl (`pip install openpyxl`).",
              file=sys.stderr)
        sys.exit(2)
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    rows_iter = ws.iter_rows(values_only=True)
    try:
        headers = ["" if v is None else str(v) for v in next(rows_iter)]
    except StopIteration:
        return [], []
    rows = []
    for r in rows_iter:
        row = ["" if v is None else str(v) for v in r]
        if any(c.strip() for c in row):
            rows.append(row)
    return headers, rows


def parse_txt_free_text(text):
    """Free text: moi block cach nhau boi dong trong = 1 case.
    Dong dau = title, cac dong sau la steps, dong bat dau 'Expected'/
    'Ket qua' la expected."""
    blocks = re.split(r"\n\s*\n", text.strip())
    cases = []
    for block in blocks:
        lines = [l.strip() for l in block.splitlines() if l.strip()]
        if not lines:
            continue
        title = lines[0]
        steps, expected = [], ""
        for l in lines[1:]:
            m = re.match(r"^(expected|ket qua|kq)\s*[:.]\s*(.*)$", l, re.I)
            if m:
                expected = (expected + "\n" + m.group(2)).strip()
            else:
                steps.append(l)
        cases.append({
            "id": "", "title": title, "priority": "", "tags": "",
            "preconditions": "", "steps": steps, "expected": expected,
            "data": "", "depends_on": "",
        })
    seen_ids = set()
    counter = [0]
    return [make_case(c, {}, seen_ids, counter) for c in cases]


def parse_json_or_yaml(text, is_yaml=False):
    if is_yaml:
        try:
            import yaml
        except ImportError:
            print("Loi: doc file .yaml can cai PyYAML (`pip install pyyaml`).",
                  file=sys.stderr)
            sys.exit(2)
        data = yaml.safe_load(text)
    else:
        data = json.loads(text)

    if isinstance(data, dict) and "testcases" in data:
        raw_cases = data["testcases"]
    elif isinstance(data, list):
        raw_cases = data
    else:
        raw_cases = [data]

    seen_ids = set()
    counter = [0]
    cases = []
    for rc in raw_cases:
        rc = dict(rc)
        extra = {k: v for k, v in rc.items() if k not in ALIASES.values()
                  and k not in ("id", "title", "steps", "expected", "priority",
                                "tags", "data", "depends_on", "preconditions")}
        cases.append(make_case(rc, extra, seen_ids, counter))
    return cases


# ---------------------------------------------------------------------------
# Main
# ---------------------------------------------------------------------------

def load_input(args):
    if args.input == "-":
        text = sys.stdin.read()
        ext = ".txt"
    else:
        ext = "." + args.input.rsplit(".", 1)[-1].lower() if "." in args.input else ""
        if ext == ".xlsx" or ext == ".xlsm":
            headers, rows = parse_xlsx(args.input, args.sheet)
            return parse_rows(rows, headers)
        with open(args.input, "r", encoding="utf-8") as f:
            text = f.read()

    if ext == ".md":
        headers, rows = parse_markdown_table(text)
        if headers:
            return parse_rows(rows, headers)
        return parse_txt_free_text(text)
    if ext in (".csv", ".tsv"):
        delim = args.delimiter or ("\t" if ext == ".tsv" else ",")
        headers, rows = parse_csv_like(text, delim)
        return parse_rows(rows, headers)
    if ext == ".json":
        return parse_json_or_yaml(text, is_yaml=False)
    if ext in (".yaml", ".yml"):
        return parse_json_or_yaml(text, is_yaml=True)
    if ext == ".txt" or ext == "":
        stripped = text.strip()
        if stripped.startswith("{") or stripped.startswith("["):
            try:
                return parse_json_or_yaml(text, is_yaml=False)
            except json.JSONDecodeError:
                pass
        if "|" in text and text.count("|") > 3:
            headers, rows = parse_markdown_table(text)
            if headers:
                return parse_rows(rows, headers)
        return parse_txt_free_text(text)
    # fallback: thu markdown table, roi free text
    headers, rows = parse_markdown_table(text)
    if headers:
        return parse_rows(rows, headers)
    return parse_txt_free_text(text)


def print_stats(cases):
    total = len(cases)
    by_priority = {}
    flagged = 0
    for c in cases:
        by_priority[c["priority"]] = by_priority.get(c["priority"], 0) + 1
        if c.get("needs_clarification"):
            flagged += 1
    print(f"Tong so case: {total}")
    for p in sorted(by_priority):
        print(f"  {p}: {by_priority[p]}")
    print(f"Case can lam ro (needs_clarification): {flagged}")
    if flagged:
        for c in cases:
            if c.get("needs_clarification"):
                print(f"  - {c['id']}: {', '.join(c['needs_clarification'])}")


def main():
    ap = argparse.ArgumentParser(
        description="Chuan hoa test case (md/xlsx/csv/tsv/json/yaml/txt/stdin) "
                    "thanh testcases.json.")
    ap.add_argument("input", help='Duong dan file input, hoac "-" cho stdin.')
    ap.add_argument("-o", "--output", help="File testcases.json output. "
                     "Mac dinh: in ra stdout.")
    ap.add_argument("--stats", action="store_true",
                     help="Chi in bang tom tat, khong dump JSON.")
    ap.add_argument("--sheet", help="Ten sheet can doc (chi ap dung cho .xlsx).")
    ap.add_argument("--delimiter", help="Ky tu phan cach cho .csv/.tsv "
                     "(mac dinh: , cho csv, tab cho tsv).")
    args = ap.parse_args()

    cases = load_input(args)
    result = {
        "meta": {
            "source": args.input,
            "generated_at": datetime.now(timezone.utc).isoformat(),
            "count": len(cases),
        },
        "testcases": cases,
    }

    if args.stats:
        print_stats(cases)
        return

    text = json.dumps(result, ensure_ascii=False, indent=2)
    if args.output:
        with open(args.output, "w", encoding="utf-8") as f:
            f.write(text)
        print(f"Da ghi {len(cases)} case vao {args.output}")
    else:
        print(text)


if __name__ == "__main__":
    main()
